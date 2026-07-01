from __future__ import annotations

from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

try:
    from markitdown import MarkItDown
except ImportError:  # pragma: no cover - optional dependency
    MarkItDown = None


def extract_markitdown_text(path: Path) -> str | None:
    if MarkItDown is None:
        return None

    converter = MarkItDown()
    result = converter.convert(str(path))
    text = getattr(result, "text_content", None)
    if not text or not text.strip():
        return None
    return text.strip()


def extract_text_from_file(path: str, material_type: str) -> str:
    file_path = Path(path)
    markitdown_text = extract_markitdown_text(file_path)
    if markitdown_text:
        return markitdown_text

    if material_type == "pdf":
        return extract_pdf_text(file_path)
    if material_type == "docx":
        return extract_docx_text(file_path)
    if material_type == "pptx":
        return extract_pptx_text(file_path)
    raise ValueError(f"Unsupported file material type: {material_type}")


def extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[Page {index}]\n{text.strip()}")
    return "\n\n".join(pages)


def extract_docx_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    return "\n\n".join(paragraph for paragraph in paragraphs if paragraph)


def extract_pptx_text(path: Path) -> str:
    presentation = Presentation(str(path))
    slides: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            slides.append(f"[Slide {slide_index}]\n" + "\n".join(texts))
    return "\n\n".join(slides)
